import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from src.batch import read_input, read_inputs, video_target_for_row
from src.task_input import parse_upload_drama_filename


class InputTests(unittest.TestCase):
    def test_video_targets_are_grouped_by_safe_drama_name(self):
        root = Path("C:/output")
        row = {"video_id": "123", "source": {"剧名": "第一部：测试/剧"}}

        self.assertEqual(
            video_target_for_row(root, row),
            root / "videos" / "第一部：测试_剧" / "123.mp4",
        )

    def test_existing_flat_video_path_is_reused_after_folder_upgrade(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            legacy = root / "videos" / "123.mp4"
            legacy.parent.mkdir(parents=True)
            legacy.write_bytes(b"existing")
            row = {"video_id": "123", "source": {"剧名": "第一部"}}

            self.assertEqual(video_target_for_row(root, row), legacy)
    def test_filename_drama_info_extracts_platform_id_and_name(self):
        self.assertEqual(
            parse_upload_drama_filename("41000329324-垃圾桶里捡到爹.xlsx"),
            {"drama_platform_id": "41000329324", "drama_name": "垃圾桶里捡到爹"},
        )
        self.assertEqual(parse_upload_drama_filename("普通需求表.xlsx"), {})

    def test_batch_excel_import_keeps_each_filename_drama_for_video_folders(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            paths = []
            for platform_id, name, video_id in [
                ("41000329324", "垃圾桶里捡到爹", "7680000000000000001"),
                ("41000321637", "藏在保温服里的天才", "7680000000000000002"),
            ]:
                path = root / f"{platform_id}-{name}.xlsx"
                workbook = Workbook()
                workbook.active.append(["剧名", "视频ID", "热度", "点赞数", "创建时间", "原始链接"])
                # A blank drama column should inherit the workbook filename.
                workbook.active.append(["", video_id, 1, 2, "", f"https://www.douyin.com/video/{video_id}"])
                workbook.save(path)
                paths.append(path)

            rows = read_inputs(paths)

            self.assertEqual([row["source"]["剧名"] for row in rows], ["垃圾桶里捡到爹", "藏在保温服里的天才"])
            self.assertEqual(
                [video_target_for_row(root / "output", row) for row in rows],
                [
                    root / "output" / "videos" / "垃圾桶里捡到爹" / "7680000000000000001.mp4",
                    root / "output" / "videos" / "藏在保温服里的天才" / "7680000000000000002.mp4",
                ],
            )

    def test_real_input_with_incorrect_dimensions(self):
        path = Path(__file__).resolve().parents[1] / "懂小剧素材_婚房门后的秘密.xlsx"
        self.assertEqual(len(read_input(path)), 200)

    def test_preserves_duplicate_rows_and_rejects_mismatched_id(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            w = Workbook()
            w.active.append(["剧名", "视频ID", "热度", "点赞数", "创建时间", "原始链接"])
            for id_, link in [("7681538825608236331", "7681538825608236331"),
                              ("7681538825608236331", "7681538825608236331"), ("123", "456")]:
                w.active.append(["测试", id_, 1, 2, "", "https://www.douyin.com/video/" + link])
            w.save(path)
            rows = read_input(path)
            self.assertEqual(len(rows), 3)
            self.assertEqual(rows[0]["video_id"], rows[1]["video_id"])
            self.assertTrue(rows[2]["input_error"])

    def test_numeric_large_excel_id_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            w = Workbook()
            w.active.append(["剧名", "视频ID", "热度", "点赞数", "创建时间", "原始链接"])
            w.active.append(["测试", 7681538825608236331, 1, 2, "", "https://www.douyin.com/video/7681538825608236331"])
            w.save(path)
            self.assertTrue(read_input(path)[0]["input_error"])
