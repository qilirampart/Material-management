import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from src.batch import run_batch


class RecoveryTests(unittest.TestCase):
    def make_input(self, folder):
        path = Path(folder) / "input.xlsx"
        workbook = Workbook()
        workbook.active.append(["剧名", "视频ID", "热度", "点赞数", "创建时间", "原始链接"])
        workbook.active.append(["=danger()", "123", 1, 2, "", "https://www.douyin.com/video/123"])
        workbook.save(path)
        return path

    def test_resume_skips_review_but_changed_rule_reruns(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.make_input(folder)
            video = Path(folder) / "out" / "videos" / "123.mp4"
            video.parent.mkdir(parents=True)
            video.write_bytes(b"test video")
            image = Path(folder) / "frame.png"
            image.write_bytes(b"test image")
            config = {"logo_reference": str(image)}
            frames = [{"path": str(image), "seconds": i} for i in range(3)]
            with patch("src.batch.Downloader"), patch("src.batch.media.validate_video", return_value={"audio": True}), \
                 patch("src.batch.vision.load_profile", return_value={}), patch("src.batch.vision.fingerprint", return_value="v1") as fp, \
                 patch("src.batch.media.extract_frames", return_value=frames), \
                 patch("src.batch.vision.review", return_value={"status": "sample_clear", "hits": []}) as model:
                run_batch(path, video.parents[1], config)
                run_batch(path, video.parents[1], config, resume=True)
                self.assertEqual(model.call_count, 1)
                fp.return_value = "v2"
                run_batch(path, video.parents[1], config, resume=True)
                self.assertEqual(model.call_count, 2)

    def test_model_failure_does_not_clear_and_export_preserves_state(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.make_input(folder)
            video = Path(folder) / "out" / "videos" / "123.mp4"
            video.parent.mkdir(parents=True)
            video.write_bytes(b"test")
            with patch("src.batch.Downloader"), patch("src.batch.media.validate_video", return_value={"audio": True}), \
                 patch("src.batch.vision.load_profile", return_value={}), patch("src.batch.vision.fingerprint", return_value="v1"), \
                 patch("src.batch.media.extract_frames", return_value=[]), patch("src.batch.vision.review", side_effect=RuntimeError("timeout")):
                state = run_batch(path, video.parents[1], {"logo_reference": str(video)})
                record = state["records"]["123"]
                self.assertEqual(record["status"], "review_required")
                self.assertFalse(record["uploaded"])
                w = load_workbook(video.parents[1] / "检测结果.xlsx")
                self.assertEqual(w.active["J2"].value, "待复核")
                w.close()
                saved = json.loads((video.parents[1] / "results.json").read_text(encoding="utf-8"))
                self.assertEqual(saved["records"]["123"]["status"], "review_required")
