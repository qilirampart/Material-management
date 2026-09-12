from __future__ import annotations

import hashlib
import json
import re
from concurrent.futures import FIRST_COMPLETED, ProcessPoolExecutor, wait
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src import media, vision
from src.download_support import Downloader

HEADERS = ["剧名", "视频ID", "热度", "点赞数", "创建时间", "原始链接"]
LABELS = {"pending": "待处理", "download_failed": "下载失败", "review_required": "待复核",
          "blocked": "命中拦截", "sample_clear": "片头抽检未发现", "invalid_input": "输入无效"}


def error_summary(exc):
    text = re.sub(r"https?://[^\s'\"<>]+", "[地址已隐藏]", str(exc))
    text = re.sub(r"(?i)(api[_-]?key|token|authorization|cookie)[=:][^\s,;]+", r"\1=[隐藏]", text)
    return f"{type(exc).__name__}: {text[:1500]}"


def download_concurrency_for(config, operation):
    """Return a conservative process count for browser-backed downloads."""
    if operation not in {"download", "both"}:
        return 1
    try:
        requested = int(config.get("download_concurrency", 1))
    except (TypeError, ValueError):
        requested = 1
    # Each worker owns a Qt WebEngine probe. More than three competing probes
    # increases failure risk and did not improve the measured wall-clock time.
    return max(1, min(requested, 3))


def _download_worker(url, target_text, config, allow_redownload):
    """Run one download in its own process so Qt WebEngine is never shared."""
    import logging
    # Windows process spawning starts a fresh interpreter, so the parent
    # logger settings are not inherited.
    logging.getLogger("src.vendor.douyin").setLevel(logging.CRITICAL)
    logging.getLogger("src.vendor.failover").setLevel(logging.CRITICAL)
    target = Path(target_text)
    try:
        if target.is_file():
            try:
                metadata = media.validate_video(target)
                return {"ok": True, "metadata": metadata}
            except (media.FFmpegError, ValueError):
                if not allow_redownload:
                    raise
                import uuid
                target.replace(target.with_name(f"{target.stem}.invalid-{uuid.uuid4().hex[:8]}{target.suffix}"))
        metadata = Downloader(config).download(url, target)
        return {"ok": True, "metadata": metadata}
    except Exception as exc:  # Worker results must be serializable and safe for UI logs.
        return {"ok": False, "reason": error_summary(exc)}


def read_input(path):
    if Path(path).suffix.lower() == '.json':
        rows = json.loads(Path(path).read_text(encoding='utf-8'))
        if not isinstance(rows, list):
            raise ValueError('候选输入必须是列表')
        for row in rows:
            if not re.fullmatch(r'(?:[0-9]{1,25}|(?:link|local)_[a-f0-9]{16})', row.get('video_id', '')):
                row['input_error'] = '素材 ID 无效'
        return rows
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = []
    try:
        for sheet in workbook:
            sheet.reset_dimensions()
            values = sheet.iter_rows(values_only=True)
            header = next(values, ())
            if not any(header):
                continue
            names = [str(x or "").strip() for x in header]
            if not all(x in names for x in HEADERS):
                raise ValueError(f"工作表 {sheet.title} 缺少必要表头：{HEADERS}")
            for row_number, raw in enumerate(values, 2):
                if not any(x is not None for x in raw):
                    continue
                row = {name: raw[names.index(name)] if names.index(name) < len(raw) else None for name in HEADERS}
                id_ = str(row["视频ID"] or "").strip()
                url = str(row["原始链接"] or "").strip()
                parsed = urlsplit(url)
                error = ""
                if not isinstance(row["视频ID"], str) or not re.fullmatch(r"\d{1,25}", id_):
                    error = "视频ID须为文本数字，防止Excel丢失精度"
                elif parsed.scheme != "https" or parsed.hostname not in {"www.douyin.com", "douyin.com"} or parsed.path != f"/video/{id_}":
                    error = "原始链接必须是与视频ID一致的抖音视频HTTPS链接"
                rows.append({"sheet": sheet.title, "row": row_number, "source": row,
                             "video_id": id_, "url": url, "input_error": error})
    finally:
        workbook.close()
    return rows


def save_json(path, payload):
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    temporary.replace(path)


def file_hash(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def export_report(rows, records, path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "检测结果"
    extra = [
        "原工作表", "原行号", "下载状态", "抽检状态", "命中项", "证据说明",
        "本地视频", "证据帧", "失败或待复核原因", "检测范围", "模型", "上传状态",
        "分辨率", "视频编码", "帧率", "总码率(Mbps)", "视频码率(Mbps)", "文件大小(MB)", "时长(s)",
    ]
    sheet.append(HEADERS + extra)
    for row in rows:
        r = records.get(row["video_id"], {}) if not row["input_error"] else {"status": "invalid_input", "reason": row["input_error"]}
        frames = r.get("frames", [])
        hits = r.get("hits", [])
        metadata = r.get("metadata", {})
        evidence = "\n".join(f"帧{h['frame']}：{h['evidence']}" for h in hits)
        sheet.append([row["source"][h] for h in HEADERS] + [row["sheet"], row["row"], r.get("download", "未下载"),
                     LABELS.get(r.get("status", "pending"), "待复核"), "、".join(dict.fromkeys(h["match"] for h in hits)), evidence,
                     r.get("video_path", ""), "\n".join(f"{f['seconds']:.3f}s {f['path']}" for f in frames), r.get("reason", ""),
                     "仅片头三帧，不代表全片", r.get("model", ""), "未上传",
                     (
                         f"{metadata.get('width')}×{metadata.get('height')}"
                         if metadata.get("width") and metadata.get("height") else ""
                     ),
                     metadata.get("video_codec", ""),
                     metadata.get("frame_rate", ""),
                     round(float(metadata.get("total_bitrate_bps", 0)) / 1_000_000, 3)
                     if metadata.get("total_bitrate_bps") else "",
                     round(float(metadata.get("video_bitrate_bps", 0)) / 1_000_000, 3)
                     if metadata.get("video_bitrate_bps") else "",
                     round(float(metadata.get("file_size_bytes", 0)) / 1048576, 2)
                     if metadata.get("file_size_bytes") else "",
                     round(float(metadata.get("duration", 0)), 3)
                     if metadata.get("duration") else "",
                     ])
        for cell in sheet[sheet.max_row]:
            if cell.data_type == "f":
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        color = {"blocked": "FDE9E7", "review_required": "FFF2CC", "sample_clear": "E2F0D9"}.get(r.get("status"))
        if color:
            sheet.cell(sheet.max_row, 10).fill = PatternFill("solid", fgColor=color)
    for cell in sheet[1]:
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="24476B")
    widths = [
        24, 24, 10, 10, 22, 48, 20, 10, 14, 24, 30, 45, 55, 60, 45, 30, 24, 12,
        16, 14, 12, 18, 18, 16, 14,
    ]
    from openpyxl.utils import get_column_letter
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    workbook.save(tmp)
    tmp.replace(path)


def run_batch(input_path, output, config, limit=None, resume=False, download_only=False,
              *, on_event=None, should_stop=None, selected_ids=None, force_review=False, operation=None):
    operation = operation or ('download' if download_only else 'both')
    if operation not in {'download', 'detect', 'both'}:
        raise ValueError('未知处理模式')
    download_only = operation == 'download'
    rows = read_input(input_path)
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    state_path = output / "results.json"
    input_digest = file_hash(input_path)
    if state_path.exists() and not resume:
        raise ValueError("批次已存在，请使用 --resume 或新输出目录")
    state = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {"input_sha256": input_digest, "records": {}}
    if state["input_sha256"] != input_digest:
        raise ValueError("输入文件与现有批次不同，请使用新输出目录")
    records = state["records"]
    profile = None if download_only else vision.load_profile(config)
    logo = Path(config.get("logo_reference", "assets/references/hongguo_logo.png"))
    if not logo.is_absolute():
        logo = vision.ROOT / logo
    phrases = config.get("phrases", vision.PHRASES)
    stamp = "" if download_only else vision.fingerprint(profile, logo, phrases)
    downloader = None if operation == 'detect' else Downloader(config)
    # Suppress signed media URLs and parser credentials in inherited download logging.
    import logging
    logging.getLogger("src.vendor.douyin").setLevel(logging.CRITICAL)
    logging.getLogger("src.vendor.failover").setLevel(logging.CRITICAL)
    seen = set()
    selected_set = set(selected_ids) if selected_ids is not None else None
    work_ids = []
    candidate_seen = set()
    for row in rows:
        candidate_id = row["video_id"]
        if row["input_error"] or candidate_id in candidate_seen:
            continue
        if selected_set is not None and candidate_id not in selected_set:
            continue
        if limit is not None and len(work_ids) >= limit:
            break
        candidate_seen.add(candidate_id)
        work_ids.append(candidate_id)
    task_positions = {video_id: index for index, video_id in enumerate(work_ids, 1)}
    task_total = len(work_ids)

    def emit(kind, **data):
        if on_event:
            on_event({"type": kind, **data})

    def persist(report=False):
        state["updated_at"] = datetime.now().isoformat()
        state["input_rows"] = rows
        state["input_path"] = str(Path(input_path).resolve())
        save_json(state_path, state)
        emit("saved", output=str(output))
        if not report:
            return
        try:
            export_report(rows, records, output / "检测结果.xlsx")
        except PermissionError:
            print("结果Excel正被占用；JSON进度已保存，关闭Excel后使用 export 命令导出。", flush=True)

    persist()
    parallel_failures = {}
    download_concurrency = download_concurrency_for(config, operation)
    if download_concurrency > 1:
        jobs = []
        scheduled_ids = set()
        for row in rows:
            id_ = row["video_id"]
            if id_ in scheduled_ids or id_ not in task_positions or row.get("local_path"):
                continue
            target = output / "videos" / f"{id_}.mp4"
            # Existing files are validated by the normal path below. Only files
            # that need network work enter independent browser processes.
            if target.is_file():
                continue
            scheduled_ids.add(id_)
            jobs.append((id_, row["url"], str(target), task_positions[id_]))

        if jobs:
            emit("download_batch", running=0, total=len(jobs), workers=download_concurrency)
            job_iterator = iter(jobs)
            futures = {}
            paused = False

            def submit_available(executor):
                nonlocal paused
                while len(futures) < download_concurrency and not paused:
                    if should_stop and should_stop():
                        paused = True
                        break
                    try:
                        video_id, url, target, task_index = next(job_iterator)
                    except StopIteration:
                        break
                    future = executor.submit(_download_worker, url, target, config, True)
                    futures[future] = video_id
                    emit("progress", video_id=video_id, stage="并发下载中",
                         task_index=task_index, task_total=task_total,
                         running=len(futures), workers=download_concurrency)

            with ProcessPoolExecutor(max_workers=download_concurrency) as executor:
                submit_available(executor)
                while futures:
                    completed, _ = wait(futures, return_when=FIRST_COMPLETED)
                    for future in completed:
                        video_id = futures.pop(future)
                        try:
                            result = future.result()
                        except Exception as exc:
                            result = {"ok": False, "reason": error_summary(exc)}
                        if not result.get("ok"):
                            parallel_failures[video_id] = result.get("reason", "下载失败")
                    submit_available(executor)

    for row in rows:
        id_ = row["video_id"]
        if row["input_error"] or id_ in seen:
            continue
        if selected_set is not None and id_ not in selected_set:
            continue
        if should_stop and should_stop():
            persist(True)
            emit("paused")
            return state
        if limit is not None and len(seen) >= limit:
            break
        seen.add(id_)
        task_index = task_positions[id_]
        if downloader is not None:
            downloader.on_event = lambda event, video_id=id_, index=task_index: on_event({
                **event,
                'video_id': video_id,
                'task_index': index,
                'task_total': task_total,
            }) if on_event else None
        target = output / "videos" / f"{id_}.mp4"
        old = records.get(id_, {})
        if row.get('local_path'):
            target = Path(row['local_path'])
        elif operation == 'detect' and old.get('video_path'):
            target = Path(old['video_path'])
        r = {"video_id": id_, "status": "pending", "download": "未下载", "uploaded": False,
             "video_path": str(target), "reason": "", "frames": [], "hits": []}
        records[id_] = r
        if id_ in parallel_failures:
            r.update(status="download_failed", reason=f"下载或校验失败：{parallel_failures[id_]}；可用 --resume 重试")
            persist()
            print("  下载失败", flush=True)
            continue
        emit("progress", video_id=id_, stage="准备处理", task_index=task_index, task_total=task_total)
        print(f"[{len(seen)}] {id_} 下载/校验", flush=True)
        try:
            if not target.is_file() and (operation == 'detect' or row.get('local_path')):
                r.update(status='review_required', reason='缺少本地视频，请先下载或添加本地视频')
                persist()
                continue
            if target.exists():
                emit('progress', video_id=id_, stage='校验本地视频', task_index=task_index, task_total=task_total)
                try:
                    meta = media.validate_video(target)
                except (media.FFmpegError, ValueError):
                    if operation == 'detect' or row.get('local_path'):
                        raise
                    import uuid
                    quarantine = target.with_name(f"{id_}.invalid-{uuid.uuid4().hex[:8]}.mp4")
                    target.replace(quarantine)
                    meta = downloader.download(row["url"], target)
            else:
                meta = downloader.download(row["url"], target)
            digest = file_hash(target)
            r.update(download="已下载", metadata=meta, video_sha256=digest, status="review_required")
        except Exception as exc:
            # Avoid exposing request URLs/API secrets through third-party exception strings.
            r.update(status="review_required" if operation == "detect" or row.get("local_path") else "download_failed", reason=f"下载或校验失败：{error_summary(exc)}；可用 --resume 重试")
            persist()
            print("  下载失败", flush=True)
            continue
        if download_only:
            if old.get('video_sha256') == digest and old.get('status') in {'sample_clear', 'blocked'}:
                records[id_] = old
            else:
                r['reason'] = '仅下载，未调用模型' 
            persist()
            continue
        if not force_review and old.get("fingerprint") == stamp and old.get("video_sha256") == digest and old.get("status") in {"blocked", "sample_clear"}:
            if old.get("frames") and all(Path(f["path"]).is_file() and file_hash(f["path"]) == f.get("sha256") for f in old["frames"]):
                records[id_] = old
                persist()
                print("  复用已验证结果", flush=True)
                continue
        persist()
        try:
            emit("progress", video_id=id_, stage="抽取片头三帧", task_index=task_index, task_total=task_total)
            r["frames"] = media.extract_frames(target, output / "frames" / id_)
            for f in r["frames"]:
                f["sha256"] = file_hash(f["path"])
            print("  三帧已提取，调用图片模型", flush=True)
            emit("progress", video_id=id_, stage="图片识别", task_index=task_index, task_total=task_total)
            r.update(vision.review(r["frames"], profile, logo, phrases))
            r["fingerprint"] = stamp
        except Exception as exc:
            r.update(status="review_required", reason=f"抽帧或识别失败：{error_summary(exc)}")
        if not meta["audio"] and r["status"] != "blocked":
            r.update(status="review_required", reason="视频无音轨，需复核下载完整性")
        persist()
        print("  " + LABELS[r["status"]], flush=True)
    persist(True)
    emit("finished")
    return state
